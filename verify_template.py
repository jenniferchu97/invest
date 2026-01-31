#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
验证杠杆率计算模板 Excel 文件
Verify leverage calculation template Excel file
"""

from openpyxl import load_workbook
from openpyxl.styles import PatternFill

def verify_excel_template():
    """验证 Excel 模板的完整性和正确性"""
    
    print("=" * 80)
    print("开始验证杠杆率计算模板 Excel 文件")
    print("=" * 80)
    
    # 加载文件
    wb = load_workbook('/home/runner/work/invest/invest/leverage_calculation_template.xlsx')
    
    # 1. 验证 Sheet 数量和名称
    print("\n1. 验证 Sheet 数量和名称:")
    expected_sheets = ['银利3号', '银利5号', '健康险']
    actual_sheets = wb.sheetnames
    print(f"   期望的 Sheet: {expected_sheets}")
    print(f"   实际的 Sheet: {actual_sheets}")
    assert actual_sheets == expected_sheets, "Sheet 名称不匹配"
    print("   ✓ Sheet 验证通过")
    
    # 2. 验证每个 Sheet 的结构
    accounts_config = {
        '银利3号': {'name': '银利3号', 'limit': 1.40},
        '银利5号': {'name': '银利5号', 'limit': 1.40},
        '健康险': {'name': '健康险', 'limit': 1.80}
    }
    
    for sheet_name, config in accounts_config.items():
        print(f"\n2. 验证 Sheet '{sheet_name}':")
        ws = wb[sheet_name]
        
        # 验证账户名称
        assert ws['B2'].value == config['name'], f"账户名称不匹配: {ws['B2'].value}"
        print(f"   ✓ 账户名称: {ws['B2'].value}")
        
        # 验证四个部分的标题
        part1_title = ws['A1'].value
        assert "第一部分" in part1_title, "第一部分标题不正确"
        print(f"   ✓ {part1_title}")
        
        # 查找第二部分
        found_part2 = False
        for row in range(1, 50):
            if ws[f'A{row}'].value and "第二部分" in str(ws[f'A{row}'].value):
                print(f"   ✓ {ws[f'A{row}'].value}")
                found_part2 = True
                break
        assert found_part2, "未找到第二部分"
        
        # 查找第三部分
        found_part3 = False
        for row in range(1, 50):
            if ws[f'A{row}'].value and "第三部分" in str(ws[f'A{row}'].value):
                print(f"   ✓ {ws[f'A{row}'].value}")
                found_part3 = True
                break
        assert found_part3, "未找到第三部分"
        
        # 查找第四部分
        found_part4 = False
        for row in range(1, 50):
            if ws[f'A{row}'].value and "第四部分" in str(ws[f'A{row}'].value):
                print(f"   ✓ {ws[f'A{row}'].value}")
                found_part4 = True
                break
        assert found_part4, "未找到第四部分"
        
        # 3. 验证输入单元格的格式（黄色填充）
        print(f"\n3. 验证输入单元格格式:")
        # 检查数据日期单元格
        date_cell = ws['B3']
        cell_color = date_cell.fill.start_color.rgb
        # 颜色可能包含透明度通道，所以检查是否包含 FFF2CC
        assert 'FFF2CC' in str(cell_color), \
            f"数据日期单元格颜色不正确: {cell_color}"
        print(f"   ✓ 输入单元格使用黄色填充 (颜色: {cell_color})")
        
        # 4. 验证公式单元格
        print(f"\n4. 验证关键公式:")
        # 净资产公式
        net_asset_formula = ws['B8'].value
        assert net_asset_formula and '=' in net_asset_formula, "净资产公式缺失"
        print(f"   ✓ 净资产公式: {net_asset_formula}")
        
        # 基准杠杆率公式
        leverage_formula = ws['B9'].value
        assert leverage_formula and 'IFERROR' in leverage_formula, "基准杠杆率公式缺失"
        print(f"   ✓ 基准杠杆率公式: {leverage_formula}")
        
        # 5. 验证杠杆限制值
        print(f"\n5. 验证杠杆限制:")
        # 查找杠杆限制单元格
        found_limit = False
        for row in range(1, 60):
            if ws[f'A{row}'].value and "杠杆限制:" in str(ws[f'A{row}'].value):
                limit_value = ws[f'B{row}'].value
                assert abs(limit_value - config['limit']) < 0.01, \
                    f"杠杆限制值不正确: {limit_value} vs {config['limit']}"
                print(f"   ✓ 杠杆限制: {limit_value} (期望: {config['limit']})")
                found_limit = True
                break
        assert found_limit, "未找到杠杆限制"
        
        # 6. 验证预警状态公式
        print(f"\n6. 验证预警状态:")
        found_warning = False
        for row in range(1, 60):
            if ws[f'A{row}'].value and "预警状态:" in str(ws[f'A{row}'].value):
                warning_formula = ws[f'B{row}'].value
                assert warning_formula and 'IF' in warning_formula, "预警状态公式缺失"
                print(f"   ✓ 预警状态公式: {warning_formula[:50]}...")
                found_warning = True
                break
        assert found_warning, "未找到预警状态"
        
        # 7. 验证冻结窗格
        print(f"\n7. 验证其他设置:")
        assert ws.freeze_panes is not None, "冻结窗格未设置"
        print(f"   ✓ 冻结窗格: {ws.freeze_panes}")
        
        print(f"\n   Sheet '{sheet_name}' 验证完成！\n")
    
    print("=" * 80)
    print("✓ 所有验证通过！模板文件创建成功！")
    print("=" * 80)
    
    # 输出使用说明
    print("\n使用说明:")
    print("1. 打开 leverage_calculation_template.xlsx 文件")
    print("2. 在每个账户的 Sheet 中，填写以下初始数据：")
    print("   - 数据日期")
    print("   - 总资产、总负债")
    print("   - 卖出回购金融资产款、交易性金融负债-买断、交易性金融负债-借券、质押式逆回购")
    print("3. 在「当日操作输入区」填写当日的操作数据")
    print("4. 系统会自动计算：")
    print("   - 净资产和基准杠杆率")
    print("   - 变化前后的各项指标对比")
    print("   - 杠杆率预警状态")
    print("5. 如果杠杆率超过限制，预警状态会显示红色警告")
    print("\n各账户杠杆限制：")
    print("   - 银利3号: 140%")
    print("   - 银利5号: 140%")
    print("   - 健康险: 180%")


if __name__ == "__main__":
    verify_excel_template()
