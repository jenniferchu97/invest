#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
创建一个带示例数据的 Excel 杠杆率计算模板
用于演示模板的使用效果
"""

from openpyxl import load_workbook


def create_demo_template():
    """创建带示例数据的模板"""
    print("=" * 60)
    print("创建带示例数据的杠杆率计算模板")
    print("=" * 60)
    
    # 加载模板
    wb = load_workbook('leverage_calculation_template.xlsx')
    
    # 示例数据
    demo_data = {
        "银利3号": {
            "总资产": 500000,  # 50亿元（以万元为单位）
            "卖出回购金融资产款": 120000,
            "交易性金融负债（买断）": 20000,
            "交易性金融负债（借券）": 15000,
            "买入返售金融资产": 10000,
            "其他负债": 5000,
            "买入券": 10000,
            "卖出回购融资": 8000
        },
        "银利5号": {
            "总资产": 300000,  # 30亿元
            "卖出回购金融资产款": 80000,
            "交易性金融负债（买断）": 10000,
            "交易性金融负债（借券）": 8000,
            "买入返售金融资产": 5000,
            "其他负债": 2000,
            "债券借贷（借券卖空）": 5000,
            "增资": 20000
        },
        "健康险": {
            "总资产": 800000,  # 80亿元
            "卖出回购金融资产款": 280000,
            "交易性金融负债（买断）": 30000,
            "交易性金融负债（借券）": 25000,
            "买入返售金融资产": 15000,
            "其他负债": 10000,
            "买入券": 15000,
            "卖出券": 5000
        }
    }
    
    for sheet_name, data in demo_data.items():
        print(f"\n填充 {sheet_name} 示例数据...")
        ws = wb[sheet_name]
        
        # 填充初始资产负债表数据（从第5行开始）
        ws['B5'] = data.get('总资产', 0)
        ws['B6'] = data.get('卖出回购金融资产款', 0)
        ws['B7'] = data.get('交易性金融负债（买断）', 0)
        ws['B8'] = data.get('交易性金融负债（借券）', 0)
        ws['B9'] = data.get('买入返售金融资产', 0)
        ws['B10'] = data.get('其他负债', 0)
        
        # 填充操作输入数据（从第18行开始）
        ws['B18'] = data.get('买入券', 0)
        ws['B19'] = data.get('卖出券', 0)
        ws['B20'] = data.get('债券借贷（借券卖空）', 0)
        ws['B21'] = data.get('买断逆回购借券', 0)
        ws['B22'] = data.get('买断逆回购借券释放', 0)
        ws['B23'] = data.get('卖出回购融资', 0)
        ws['B24'] = data.get('卖出回购到期归还', 0)
        ws['B25'] = data.get('增资', 0)
        ws['B26'] = data.get('减资', 0)
        
        print(f"  ✓ 已填充 {sheet_name} 的示例数据")
        print(f"    - 总资产: {data.get('总资产', 0):,.0f} 万元")
        print(f"    - 卖出回购: {data.get('卖出回购金融资产款', 0):,.0f} 万元")
    
    # 保存示例模板
    output_file = "leverage_calculation_demo.xlsx"
    wb.save(output_file)
    
    print("\n" + "=" * 60)
    print(f"✓ 示例模板已生成: {output_file}")
    print("=" * 60)
    print("\n请使用 Excel 打开该文件查看：")
    print("  1. 自动计算的总负债、净资产和杠杆率")
    print("  2. 操作前后的对比数据")
    print("  3. 杠杆限制预警状态")
    print("\n注意：公式计算结果需要在 Excel 中打开文件才能正确显示")


if __name__ == "__main__":
    create_demo_template()
